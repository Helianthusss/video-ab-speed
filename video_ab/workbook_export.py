"""Portable XLSX export using the public openpyxl package."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NUM!")


def build_workbook(tabs, output_path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, table in tabs.items():
        ws = wb.create_sheet(name[:31])
        ws.sheet_view.showGridLines = False
        headers = table["headers"]
        ws.append(headers)
        for row in table["rows"]:
            ws.append(row)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="193C50")
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, header in enumerate(headers, 1):
            width = max(
                12,
                min(
                    36,
                    max(
                        [len(str(header))]
                        + [
                            len(str(r[index - 1])) if r[index - 1] is not None else 0
                            for r in table["rows"][:100]
                        ]
                    )
                    + 2,
                ),
            )
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "A2"
        if name == "Raw":
            positions = {header: i + 1 for i, header in enumerate(headers)}
            required = ["Dir.", "L (m)", "Δt", "V (km/h)", "tA_corrected", "tB_corrected"]
            if all(k in positions for k in required):
                for row_index in range(2, ws.max_row + 1):
                    d = get_column_letter(positions["Dir."])
                    length = get_column_letter(positions["L (m)"])
                    dt = get_column_letter(positions["Δt"])
                    speed = get_column_letter(positions["V (km/h)"])
                    a = get_column_letter(positions["tA_corrected"])
                    b = get_column_letter(positions["tB_corrected"])
                    ws[f"{dt}{row_index}"] = (
                        f'=IF(OR(ISBLANK({a}{row_index}),ISBLANK({b}{row_index})),"",IF({d}{row_index}="A→B",{b}{row_index}-{a}{row_index},{a}{row_index}-{b}{row_index}))'
                    )
                    ws[f"{speed}{row_index}"] = (
                        f'=IF(OR({dt}{row_index}="",{dt}{row_index}<=0,{length}{row_index}<=0),"",3.6*{length}{row_index}/{dt}{row_index})'
                    )
                    ws[f"{dt}{row_index}"].number_format = "0.000000"
                    ws[f"{speed}{row_index}"].number_format = "0.00"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    verify_workbook(output_path)


def verify_workbook(path):
    # read_only keeps the file handle open until closed, which blocks deletion on Windows.
    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        if not wb.sheetnames:
            raise ValueError("Workbook không có sheet")
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and any(
                        token in cell.value for token in ERROR_TOKENS
                    ):
                        raise ValueError(f"Lỗi workbook tại {ws.title}!{cell.coordinate}")
    finally:
        wb.close()
