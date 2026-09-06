from pathlib import Path
import csv,json,zipfile,math,os
from core import *
from media import metadata
ROOT=Path(__file__).resolve().parent
OUTPUT_DIR=Path(os.environ.get('AB_OUTPUT_DIR',ROOT/'outputs')).expanduser().resolve()
def cell(v):
 if isinstance(v,(dict,list)):return json.dumps(v,ensure_ascii=False)
 return v

def tables(s):
 d=summary(s);label=DEMO if s['mode']=='Demo' else 'RESEARCH';raw=[];sens=[]
 for r in d['records']:
  raw.append({'Site':s['site'],'ID':r['id'],'Type':r['type'],'Dir.':r['direction'],'tA':r.get('tA'),'tB':r.get('tB'),'L (m)':s['L'],'Δt':r['dt'],'V (km/h)':r['speed'],'QC':r['qc'],**{k:r.get(k) for k in ['frameA','frameB','ptsA','ptsB','time_baseA','time_baseB','tA_corrected','tB_corrected','videoA','videoB','observer','protocol_version','sampling_status','matching','decision','reason','reviewer','created','updated','warnings','duplicate_event','relative_uncertainty']},'session':s['id'],'interval':r['interval'],'correction_version':s['sync']['version'],'Mode':label})
  if r.get('frameA') is not None and r.get('frameB') is not None:
   ma=metadata(r['videoA']);mb=metadata(r['videoB'])
   for da in [-1,0,1]:
    for db in [-1,0,1]:
     ia=r['frameA']+da;ib=r['frameB']+db
     if 0<=ia<len(ma['frames']) and 0<=ib<len(mb['frames']):
      x=dict(r,tA=ma['frames'][ia]['time'],tB=mb['frames'][ib]['time']);sens.append(dict(ID=r['id'],frame_delta_A=da,frame_delta_B=db,speed=measure(x,s)['speed'],scenario='actual adjacent PTS',Mode=label))
   for off in s['sync'].get('observed_offset_delta',[]):
    import copy
    ss=copy.deepcopy(s);ss['sync']['offset_start']+=off;ss['sync']['offset_end']+=off
    sens.append(dict(ID=r['id'],offset_delta=off,speed=measure(r,ss)['speed'],scenario='observed offset perturbation',Mode=label))
 summary_rows=[];class_rows=[]
 for x in d['rows']:
  z=x['stats'];mc=x['classes']['MC'];car=x['classes']['Car']
  summary_rows.append(dict(zip(['Site','Dir.','N15','ns','MC%','Mean','95% CI','Median','V85','SD','CV','nMC','Mean MC','nCar','Mean Car','Diff.'],[s['site'],x['direction'],x['N15'],x['ns'],x['MC_percent'],z['Mean'],z['CI'],z['Median'],z['V85'],z['SD'],z['CV'],mc['n'],mc['Mean'],car['n'],car['Mean'],x['Diff']]))|{'Diff_abs':x['Diff_abs'],'session':s['id'],'interval':x['interval'],'T_seconds':x['duration'],'N_observed':x['N_observed'],'q_equivalent':x['q_equivalent'],'selected':x['selected'],'matched':x['matched'],'source_ids':x['source_ids'],'scope':x['scope'],'Mode':label})
  for c,cs in x['classes'].items():class_rows.append(dict(Site=s['site'],direction=x['direction'],interval=x['interval'],Type=c,**cs,source_ids=[r['id'] for r in d['records'] if r['valid'] and r['type']==c and r['direction']==x['direction'] and r['interval']==x['interval']],Mode=label))
 audits=[];pairs=[]
 for a in s['audits']:
  r=next(r for r in d['records'] if r['id']==a['record_id']);out=dict(a)
  if a.get('status')=='done':
   rep=dict(r,tA=a['tA'],tB=a['tB']);v=measure(rep,s)['speed'];out.update(reference_speed=r['speed'],repeat_speed=v,frame_difference_A=a['frameA']-r['frameA'],frame_difference_B=a['frameB']-r['frameB'],time_difference_A=a['tA']-r['tA'],time_difference_B=a['tB']-r['tB'])
   if r['speed'] is not None and v is not None:pairs.append((r['speed'],v))
  audits.append(out)
 metrics=agreement(pairs);done=[a for a in s['audits'] if a['status']=='done'];metrics.update(matching_denominator=len(done),correct=sum(a['matching']=='correct' for a in done),incorrect=sum(a['matching']=='incorrect' for a in done),uncertain=sum(a['matching']=='uncertain' for a in done),claim='Reliability; không chứng minh độ chính xác ngoài hiện trường')
 flow=[dict(f,T_seconds=min(s['interval_seconds'],s['end']-s['start']-f['interval']*s['interval_seconds']),Mode=label) for f in s['flow']]
 tabs={'Raw':raw,'Site summary':summary_rows,'Class summary':class_rows,'Setup':[dict(field=k,value=cell(s[k]),source='Nhập tay' if s[k] not in [None,'',{}] else 'Chưa có dữ liệu') for k in ['project','site','location','observer','session_start','timezone','reference','L','uL','uT','start','end','interval_seconds','setup']], 'Site changes':s['site_changes'],'Events':s['events'],'Video sync':[{'camera':c,**{k:v for k,v in metadata(s['video'+c]).items() if k not in ['frames','path']},'sync':s['sync']} for c in ['A','B'] if s['video'+c]],'Flow count':flow,'Sampling':s['plans'],'QC exclusion':[r for r in raw if r['QC']!=[0] or r['decision']!='keep'],'Audit':audits,'Reliability':[{'metric':k,'value':cell(v)} for k,v in metrics.items()],'Sensitivity':sens,'Protocol':[{'field':k,'value':cell(v)} for k,v in s['protocol'].items()],'Sources':s.get('sources',[])}
 defaults={'Raw':['Site','ID','Type','Dir.','tA','tB','L (m)','Δt','V (km/h)','QC'],'Site changes':['feature','before','after','evidence','created','observer'],'Events':['feature','before','after','evidence','created','observer'],'Flow count':['id','direction','section','interval','type','count','source'],'Sampling':['N','n_target','k','start','seed','selected'],'Audit':['id','record_id','kind','observer','status','frameA','frameB'],'QC exclusion':['ID','QC','decision','reason','reviewer'],'Sensitivity':['ID','scenario','speed'],'Sources':['title','url','author','license','accessed','segment','assumptions']}
 computed={'Δt','V (km/h)','Mean','95% CI','Median','V85','SD','CV','Diff.','Diff_abs','MC%','ns','nMC','nCar','Mean MC','Mean Car','tA_corrected','tB_corrected','q_equivalent','P85_CI','relative_uncertainty','speed'}
 decoded={'tA','tB','ptsA','ptsB','time_baseA','time_baseB','width','height','fps','time_base','timing','first','last'}
 clicked={'frameA','frameB','ID','created','updated'}
 dictionary=[];out={}
 for name,rows in tabs.items():
  headers=list(dict.fromkeys([k for r in rows for k in r])) or defaults.get(name,['field','value','source'])
  if 'Mode' not in headers:headers.append('Mode')
  out[name]={'headers':headers,'rows':[[cell(r.get(k,label if k=='Mode' else None)) for k in headers] for r in rows]}
  for k in headers:
   source='Tự tính' if k in computed else 'Đọc từ video' if k in decoded else 'Sinh từ thao tác click' if k in clicked else 'Nhập tay'
   if not rows or all(r.get(k) is None for r in rows):source='Chưa có dữ liệu'
   dictionary.append(dict(form=name,field=k,source=source,unit='s' if k in ['tA','tB','Δt','tA_corrected','tB_corrected','T_seconds'] else 'km/h' if k in ['V (km/h)','Mean','SD','Median','V85','Mean MC','Mean Car','Diff.','Diff_abs'] else None,definition='0-based frame / interval' if k in ['frameA','frameB','interval'] else k,Mode=label))
 out['Dictionary']={'headers':list(dictionary[0]),'rows':[list(r.values()) for r in dictionary]}
 return out,d,metrics

def export_bundle(s):
 out=OUTPUT_DIR/s['id'];out.mkdir(parents=True,exist_ok=True);tabs,d,metrics=tables(s);label=DEMO if s['mode']=='Demo' else 'RESEARCH'
 (out/'tables.json').write_text(json.dumps({'label':label,'tables':tabs},ensure_ascii=False,allow_nan=False))
 (out/'snapshot.json').write_text(json.dumps(s,ensure_ascii=False,indent=2))
 for name,table in tabs.items():
  with (out/(name.replace(' ','_')+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
   w=csv.writer(f);w.writerow(table['headers']);w.writerows(table['rows'])
 from workbook_export import build_workbook
 build_workbook(tabs,out/'Survey.xlsx')
 from docx import Document
 from docx.shared import Pt,Inches,RGBColor
 from docx.enum.text import WD_LINE_SPACING
 doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.55)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  doc.styles[name].font.name='Arial';doc.styles[name].font.color.rgb=RGBColor(0,0,0)
 doc.styles['Normal'].font.size=Pt(9.5)
 doc.styles['Normal'].paragraph_format.space_after=Pt(3)
 doc.styles['Normal'].paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE
 sec.header.paragraphs[0].text=label
 title=doc.add_paragraph()
 title.paragraph_format.space_after=Pt(10)
 run=title.add_run('Báo cáo khảo sát tốc độ trên đoạn A B')
 run.bold=True;run.font.name='Arial';run.font.size=Pt(21);run.font.color.rgb=RGBColor(0,0,0)
 doc.add_paragraph(f"Site {s['site']} | Phiên {s['id']} | {s['timezone']} | {s['start']}–{s['end']} giây")
 doc.add_paragraph('Kết quả là tốc độ hành trình trên đoạn. Các bảng đầy đủ và biểu mẫu vận hành nằm trong workbook và CSV cùng gói xuất; snapshot JSON bảo toàn dữ liệu và protocol để tái lập.')
 doc.add_paragraph('[CHƯA ĐỦ CƠ SỞ DỮ LIỆU ĐỂ KẾT LUẬN] về độ chính xác ngoài hiện trường hoặc quan hệ nhân quả với tai nạn lịch sử.')
 for row in d['rows']:
  doc.add_heading(f"Hướng {row['direction']} khoảng {row['interval']}",1)
  z=row['stats'];fmt=lambda v:'Chưa có / không tính' if v is None else str(round(v,4)) if isinstance(v,float) else str(v)
  doc.add_paragraph(f"Thời lượng thực {row['duration']} s; số đếm toàn bộ {fmt(row['N_observed'])}; xe khớp {row['matched']}; xe giữ phân tích {row['ns']}. Lưu lượng giờ tương đương {fmt(row['q_equivalent'])} xe/h; không phải AADT.")
  doc.add_paragraph('Mean '+fmt(z['Mean'])+' km/h; SD '+fmt(z['SD'])+' km/h; CV '+fmt(z['CV'])+'; Median '+fmt(z['Median'])+'; P85 '+fmt(z['V85'])+'.')
  doc.add_paragraph('95% CI Mean '+fmt(z['CI'])+'; bootstrap CI P85 '+fmt(z['P85_CI'])+'. '+row['scope']+'.')
  doc.add_paragraph('ID nguồn: '+(', '.join(row['source_ids']) or 'Chưa có'))
 doc.add_heading('Phương pháp và giới hạn',1)
 for text in ['Hiệu chỉnh B bằng cộng offset tuyến tính theo timestamp B; đồng hồ A là tham chiếu. Phiên bản '+str(s['sync']['version'])+'.','Phân vị type 7; bootstrap percentile '+str(s['protocol']['bootstrap'])+' lần với seed '+str(s['protocol']['seed'])+'. CI thường giả định độc lập và tính đại diện phù hợp. Tắt CI nếu chưa xác nhận giả định; chưa triển khai bootstrap theo khối.','n=0: chỉ n=0, chỉ tiêu khác trống. n=1: Mean và Median, không SD, CV, CI hoặc P85. Nhóm hiếm cần diễn giải mô tả; không có ngưỡng đủ mẫu P85 được xác nhận.','Precision cỡ mẫu Mean không phải sai số thiết bị và không chứng minh đủ mẫu P85. Diff. có dấu; Diff_abs là độ lớn.','Reliability: '+json.dumps(metrics,ensure_ascii=False),'Thiếu thông tin thực địa giữ trống. Bộ demo và thử công thức không kiểm chứng độ chính xác ngoài hiện trường.']:
  doc.add_paragraph(text)
 doc.save(out/'Report.docx')
 # PDF uses same report paragraphs, Unicode font, automatic page flow.
 from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
 from reportlab.lib.styles import getSampleStyleSheet
 from reportlab.pdfbase import pdfmetrics
 from reportlab.pdfbase.ttfonts import TTFont
 from xml.sax.saxutils import escape
 font_candidates=[
  Path('/System/Library/Fonts/Supplemental/Arial.ttf'),
  Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
  Path('C:/Windows/Fonts/arial.ttf'),
 ]
 font=next((candidate for candidate in font_candidates if candidate.exists()),None)
 if font is None:raise RuntimeError('Cần Arial hoặc DejaVu Sans để xuất PDF tiếng Việt')
 pdfmetrics.registerFont(TTFont('ArialUnicode',str(font)))
 styles=getSampleStyleSheet()
 for st in styles.byName.values():st.fontName='ArialUnicode';st.fontSize=10;st.leading=14
 story=[]
 for p in doc.paragraphs:
  style=styles['Heading1'] if p.style.name in ['Title','Heading 1'] else styles['Normal'];story.extend([Paragraph(escape(p.text),style),Spacer(1,8)])
 def footer(c,doc):c.setFont('ArialUnicode',7);c.drawString(35,20,label)
 SimpleDocTemplate(str(out/'Report.pdf'),rightMargin=40,leftMargin=40,topMargin=40,bottomMargin=40).build(story,onFirstPage=footer,onLaterPages=footer)
 archive=out.parent/(s['id']+'.zip')
 with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
  for f in out.iterdir():
   if f.is_file():z.write(f,f.name)
 return archive
